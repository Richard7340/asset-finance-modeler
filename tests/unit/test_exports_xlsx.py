import pytest
from openpyxl import load_workbook

from asset_finance_modeler.assets.saas.loader import load_preset
from asset_finance_modeler.assets.saas.model import SaasModel
from asset_finance_modeler.store.exports import to_xlsx


@pytest.fixture(scope="module")
def results():
    return SaasModel(load_preset("gestnova")).run()


def test_to_xlsx_creates_file(results, tmp_path):
    out = tmp_path / "model.xlsx"
    to_xlsx(results, path=str(out))
    assert out.exists()


def test_to_xlsx_has_expected_sheets(results, tmp_path):
    out = tmp_path / "model.xlsx"
    to_xlsx(results, path=str(out))
    wb = load_workbook(out)
    sheet_names = set(wb.sheetnames)
    assert {"Summary", "PnL", "CashFlow", "Balance", "UnitEcon"}.issubset(sheet_names)


def test_to_xlsx_pnl_has_data(results, tmp_path):
    out = tmp_path / "model.xlsx"
    to_xlsx(results, path=str(out))
    wb = load_workbook(out)
    pnl = wb["PnL"]
    # Header row + 60 data rows
    assert pnl.max_row == 61
    # Column A header = "period", first data col after header
    assert pnl.cell(row=1, column=1).value == "period"
    assert pnl.cell(row=2, column=1).value == 0


def test_to_xlsx_summary_has_key_metrics(results, tmp_path):
    out = tmp_path / "model.xlsx"
    to_xlsx(results, path=str(out))
    wb = load_workbook(out)
    summary = wb["Summary"]
    # Summary is 2-column key:value
    keys = [summary.cell(row=r, column=1).value for r in range(2, summary.max_row + 1)]
    assert "revenue_y1" in keys
    assert "enterprise_value" in keys
